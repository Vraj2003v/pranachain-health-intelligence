"""Build the weekly Evidence Trail workbook from log.json.

Each Sunday: edit log.json (add rows, change results), then run
    python docs/progress_log/build_log.py
The workbook is written next to this file as PranaPulse_Project_Progress_Log.xlsx.
Same layout as the Evidence Trail format: team log on top, individual contribution logs below.
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).parent
data = json.loads((HERE / "log.json").read_text(encoding="utf-8"))

NAVY, BLUE = "1F3864", "D9E2F3"
thin = Side(style="thin", color="000000")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
wb = Workbook()
ws = wb.active
ws.title = "Team Project Log"
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
for col, w in zip("ABCDEFG", [7, 13, 28, 42, 36, 36, 32]):
    ws.column_dimensions[col].width = w


def banner(row, text, size=14):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    c = ws.cell(row, 1, text)
    c.font = F(size=size, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 30


def est_height(*cells):
    lines = 1
    for text, width in cells:
        t = str(text)
        lines = max(lines, sum(max(1, -(-len(p) // max(1, int(width * 1.05)))) for p in t.split("\n")))
    return max(24, 13.5 * lines + 6)


banner(1, data["title"])
for r, key, bold in [(2, "subtitle", True), (3, "team", False)]:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(r, 1, data[key]); c.font = F(bold=bold); c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 18

RANGE = {"1": "Sep 28 - Oct 4", "2-3": "Oct 5 - Oct 18", "4": "Oct 19 - Oct 25"}
def wk(w):
    w = str(w)
    return f"Week {w}\n{RANGE[w]}" if w in RANGE else f"Week {w}"

heads = ["Week", "Week ending", "What we planned", "What we built / completed", "Evidence", "What changed", "Next step"]
for i, h in enumerate(heads, 1):
    c = ws.cell(5, i, h); c.font = F(bold=True); c.fill = PatternFill("solid", fgColor=BLUE); c.border = box
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[5].height = 24
r = 6
for t in data["team_log"]:
    vals = [wk(t["week"]), t["ending"], t["planned"], t["built"], t["evidence"], t["changed"], t["next"]]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v); c.font = F(); c.border = box
        c.alignment = Alignment(horizontal="center" if i <= 2 else "left", vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = est_height((t["planned"], 28), (t["built"], 42), (t["evidence"], 36), (t["changed"], 36), (t["next"], 32))
    r += 1

r += 1
banner(r, "INDIVIDUAL CONTRIBUTION LOG"); r += 2
for p in data["people"]:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c = ws.cell(r, 1, f'{p["name"]}  |  {p["role"]}'); c.font = F(bold=True, color=NAVY); c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r].height = 20; r += 1
    # columns: Week | My task (B:C) | Evidence (D) | Result (E) | Next task (F:G)
    for col, h in [(1, "Week"), (2, "My task"), (4, "Evidence"), (5, "Result"), (6, "Next task")]:
        c = ws.cell(r, col, h); c.font = F(bold=True); c.fill = PatternFill("solid", fgColor=BLUE); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for a, b in [(2, 3), (6, 7)]:
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    for col in range(1, 8):
        ws.cell(r, col).border = box; ws.cell(r, col).fill = PatternFill("solid", fgColor=BLUE)
    ws.row_dimensions[r].height = 24; r += 1
    for x in p["rows"]:
        for col, v in [(1, wk(x["week"])), (2, x["task"]), (4, x["evidence"]), (5, x["result"]), (6, x["next"])]:
            c = ws.cell(r, col, v); c.font = F(); c.alignment = Alignment(horizontal="center" if col in (1, 5) else "left", vertical="top", wrap_text=True)
        for a, b in [(2, 3), (6, 7)]:
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        for col in range(1, 8):
            ws.cell(r, col).border = box
        ws.row_dimensions[r].height = est_height((x["task"], 41), (x["evidence"], 42), (x["result"], 36), (x["next"], 68))
        r += 1
    r += 1
ws.freeze_panes = "A6"
out = HERE / "PranaPulse_Project_Progress_Log.xlsx"
wb.save(out)
print("wrote", out)
